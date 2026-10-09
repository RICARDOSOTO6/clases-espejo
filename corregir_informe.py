# -*- coding: utf-8 -*-
"""
Corrige el informe técnico del usuario (hasta capítulo 2) con la información
real del proyecto «Clases Conjuntas» (Comunidad de Apoyo a Clases Espejo, CACE),
tomada de docs/readme y de la carpeta docs/. Conserva el diseño, estilos y
tablas del documento original.
"""
import copy
from docx import Document
from docx.shared import Pt
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

SRC = r"C:\Users\ricar\Downloads\Informe_Tecnico_Clases_Espejo_corregido_hasta_capitulo_2.docx"
DST = r"C:\Users\ricar\Documents\ClasesEspejo\Informe_Tecnico_Clases_Espejo_corregido.docx"

doc = Document(SRC)
TNR = "Times New Roman"

# ---------------------------------------------------------------------------
# Utilidades
# ---------------------------------------------------------------------------
def clear_runs(p):
    for r in list(p.runs):
        r._element.getparent().remove(r._element)


def replace_paragraph(needle, new_text, font=TNR, size=12, bold=None, italic=None):
    """Reemplaza el texto del primer párrafo que contiene `needle`, conservando
    el estilo y la alineación del párrafo."""
    for p in doc.paragraphs:
        if needle in p.text:
            clear_runs(p)
            run = p.add_run(new_text)
            run.font.name = font
            run.font.size = Pt(size)
            if bold is not None:
                run.font.bold = bold
            if italic is not None:
                run.font.italic = italic
            return p
    raise ValueError("No se encontró el párrafo: " + needle)


# ---------------------------------------------------------------------------
# 1) RESUMEN e INTRODUCCIÓN (contexto ITESCAM + organismo externo CACE)
# ---------------------------------------------------------------------------
replace_paragraph(
    "El presente proyecto de residencia profesional describe el desarrollo",
    "El presente proyecto de residencia profesional describe el desarrollo de la plataforma web "
    "«Clases Conjuntas», un sistema para la gestión de clases espejo desarrollado para la Comunidad "
    "de Apoyo a Clases Espejo (CACE), organismo externo que promueve la colaboración académica entre "
    "instituciones educativas. La plataforma permite que dos instituciones organicen clases espejo de "
    "principio a fin: el registro de instituciones, docentes y materias, el envío y la doble revisión "
    "de solicitudes de colaboración, la conformación de proyectos, la planificación conjunta de "
    "sesiones y el seguimiento de actividades, evidencias, reportes y evaluación de cada experiencia.",
)

replace_paragraph(
    "Para su construcción se empleó una metodología ágil",
    "Para su construcción se empleó una metodología ágil (Scrum) y una arquitectura cliente-servidor: "
    "backend en NestJS 11 con Prisma ORM y PostgreSQL, y frontend en Angular 21, comunicados mediante "
    "una API REST con autenticación JWT. El resultado es una plataforma funcional que organiza el flujo "
    "completo del proceso, con dos roles (agente de internacionalización y docente), y que facilita la "
    "trazabilidad, la supervisión institucional y la documentación de las colaboraciones académicas.",
)

replace_paragraph(
    "Las clases espejo constituyen una estrategia de colaboración",
    "Las clases espejo constituyen una estrategia de colaboración académica interinstitucional que "
    "permite que estudiantes y docentes de instituciones diferentes participen de manera conjunta en "
    "actividades académicas, utilizando herramientas digitales para establecer la comunicación y la "
    "coordinación necesarias. Esta práctica forma parte de las acciones orientadas a la "
    "internacionalización en casa; sin embargo, su gestión suele realizarse mediante herramientas "
    "independientes —correo electrónico, documentos de oficina y hojas de cálculo—, lo que dificulta "
    "mantener un seguimiento ordenado de las solicitudes, los participantes, las sesiones y las "
    "evidencias generadas.",
)

replace_paragraph(
    "Ante esta situación, se plantea el desarrollo",
    "Ante esta situación se plantea el desarrollo de una plataforma web, denominada «Clases Conjuntas», "
    "que centralice y organice las diferentes etapas de gestión de una clase espejo. El proyecto se "
    "realiza como residencia profesional en el Instituto Tecnológico Superior de Calkiní (ITESCAM) para "
    "la Comunidad de Apoyo a Clases Espejo (CACE), organismo externo que agrupa a las instituciones "
    "interesadas en esta modalidad de colaboración. La solución contempla el registro de instituciones, "
    "docentes y materias; la gestión y la doble revisión de solicitudes; la conformación de proyectos; "
    "la planificación conjunta; y el seguimiento de sesiones, actividades, evidencias, reportes y "
    "evaluación. En esta etapa del informe se documentan los antecedentes, los objetivos, el alcance y "
    "el método empleado para el análisis, diseño y desarrollo de la solución.",
)

replace_paragraph(
    "El proyecto se justifica por la necesidad de optimizar",
    "El proyecto se justifica por la necesidad de dotar a la Comunidad de Apoyo a Clases Espejo (CACE) "
    "de una herramienta que formalice y ordene la gestión de las clases espejo entre sus instituciones "
    "participantes. El uso de herramientas independientes para registrar la información y dar "
    "seguimiento a las actividades provoca que los datos se encuentren dispersos y que el estado de una "
    "solicitud o proyecto sea difícil de consultar, además de complicar la integración de las evidencias "
    "y la documentación generada. La plataforma permite concentrar esta información en un solo lugar, "
    "establecer un flujo organizado de trabajo —con una doble revisión institucional de cada propuesta— "
    "y facilitar la trazabilidad de las colaboraciones académicas.",
)

replace_paragraph(
    "Alcance del proyecto. El proyecto comprende",
    "Alcance del proyecto. El proyecto comprende el análisis, diseño, desarrollo y pruebas de una "
    "plataforma web orientada a la gestión de clases espejo para la Comunidad CACE. La solución "
    "contempla dos roles (agente de internacionalización y docente) y módulos de autenticación, "
    "instituciones, docentes, materias y asignaciones, solicitudes con doble revisión, proyectos, "
    "planificación conjunta, chat, sesiones, actividades, evidencias, reportes de clase, evaluación, "
    "cierre, calendario y recordatorios.",
)

replace_paragraph(
    "Limitaciones. En esta etapa, el proyecto se encuentra",
    "Limitaciones. En esta etapa no se contempla la recuperación automática de contraseña (se gestiona "
    "a través del agente de internacionalización), el desarrollo de una aplicación móvil nativa, la "
    "integración con sistemas académicos externos ni la automatización de las herramientas de "
    "videoconferencia. Asimismo, el alcance se limita a la gestión de clases espejo y no pretende "
    "sustituir otros sistemas o procesos académico-administrativos.",
)

replace_paragraph(
    "Objetivo general. Desarrollar un sistema web",
    "Objetivo general. Desarrollar una plataforma web para la gestión de clases espejo que permita a la "
    "Comunidad de Apoyo a Clases Espejo (CACE) y a sus instituciones participantes organizar y dar "
    "seguimiento al proceso de colaboración académica interinstitucional.",
)

replace_paragraph(
    "Implementar los módulos de gestión de instituciones",
    "• Implementar los módulos de autenticación, instituciones, docentes, materias, asignaciones, "
    "solicitudes con doble revisión, proyectos, planificación, sesiones, actividades, evidencias, "
    "reportes, evaluación, cierre, calendario y recordatorios.",
)

# ---------------------------------------------------------------------------
# 2) CAPÍTULO 2 — Materiales y métodos (datos reales del proyecto)
# ---------------------------------------------------------------------------
replace_paragraph(
    "El proyecto se desarrolló mediante una metodología ágil",
    "El proyecto se desarrolló mediante una metodología ágil basada en Scrum, organizando el trabajo en "
    "iteraciones (sprints) conforme a un plan de ocho semanas que permitió la entrega incremental de los "
    "módulos del sistema. Cada iteración consideró actividades de análisis, diseño, implementación y "
    "pruebas, y el producto se sometió a tres rondas de auditoría de calidad (código, experiencia de uso "
    "y accesibilidad). Para el control de versiones se utilizó Git —46 confirmaciones documentadas entre "
    "el 28 de agosto y el 23 de septiembre de 2026— y el seguimiento se realizó mediante un tablero de "
    "tareas. La Figura 2.1 presenta el esquema general de la metodología empleada.",
)

replace_paragraph(
    "A partir del análisis del proceso actual y de las entrevistas",
    "A partir del análisis del proceso actual y de los casos de uso de los dos roles del sistema (agente "
    "de internacionalización y docente), se identificaron los principales requerimientos funcionales. "
    "Estos requerimientos se organizaron por módulo y se priorizaron de acuerdo con su importancia para "
    "el funcionamiento de la plataforma. El Cuadro 2.1 presenta el conjunto de requerimientos funcionales "
    "considerado para el desarrollo.",
)

replace_paragraph(
    "El sistema se diseñó con una arquitectura web de tres capas",
    "El sistema se diseñó con una arquitectura web de tres capas: presentación, lógica de negocio y "
    "datos. La capa de presentación corresponde al frontend, desarrollado en Angular 21 (modo zoneless "
    "con signals), que proporciona la interfaz con la que interactúa el usuario. La capa de lógica de "
    "negocio corresponde al backend, implementado en NestJS 11 con TypeScript, responsable de procesar "
    "las solicitudes, aplicar las reglas del sistema y exponer los servicios mediante una API REST de 62 "
    "endpoints. La capa de datos está constituida por una base de datos PostgreSQL gestionada con Prisma "
    "ORM. La comunicación entre el frontend y el backend se realiza mediante la API, con autenticación "
    "basada en tokens JWT. La Figura 2.2 presenta la arquitectura general.",
)

replace_paragraph(
    "El diseño de la base de datos se realizó a partir de las entidades",
    "El diseño de la base de datos se realizó a partir de las entidades que representan los elementos del "
    "proceso de clases espejo, dando como resultado un modelo de 22 tablas. Entre ellas se encuentran "
    "usuario, docente, agente de internacionalización, institución, invitación, materia, asignación "
    "docente-materia, solicitud de clase espejo, revisión de solicitud, proyecto, docente de proyecto, "
    "planificación conjunta, reporte de planificación, mensaje, sesión, actividad, evidencia, reporte de "
    "clase conjunta, participación y evaluación. Se definieron llaves primarias y foráneas para mantener "
    "la integridad referencial y el esquema se versionó mediante 10 migraciones de Prisma. La Figura 2.3 "
    "presenta el modelo entidad-relación del sistema.",
)

replace_paragraph(
    "El backend se implementó como una API REST encargada",
    "El backend se implementó en NestJS 11 con TypeScript, organizado en 12 módulos y 10 controladores, "
    "y se encarga de exponer la API REST con 62 endpoints. En esta capa se incorporaron mecanismos para "
    "validar los datos recibidos (class-validator), controlar el acceso por rol mediante JWT y bcrypt "
    "para el manejo seguro de contraseñas, así como el envío de correos de invitación (Nodemailer) y la "
    "subida de archivos de evidencia (multer). Se aplicaron principios de separación de responsabilidades "
    "y se incluyó un filtro global de errores de Prisma. La Figura 2.4 ilustra la estructura general de "
    "la API REST.",
)

replace_paragraph(
    "El frontend se desarrolló como una interfaz web responsiva",
    "El frontend se desarrolló en Angular 21, en modo zoneless con signals y componentes standalone, y "
    "presenta la identidad visual de la comunidad CACE. Ofrece dos paneles diferenciados según el rol: "
    "el panel del agente de internacionalización (institución, docentes, materias, asignaciones, "
    "solicitudes y proyectos) y el panel del docente (solicitudes y proyectos). Incluye formularios con "
    "validación, un calendario mensual, una sección de recordatorios, la gestión de la foto de perfil y "
    "una guía interactiva de primeros pasos. La Figura 2.5 presenta los prototipos de las principales "
    "pantallas del sistema.",
)

replace_paragraph(
    "La integración consistió en conectar los componentes",
    "La integración consistió en conectar el frontend con la API del backend y la base de datos. En "
    "producción, el frontend se compila y se sirve desde el propio backend (backend/public), de modo que "
    "toda la aplicación corre en un único puerto (3000) sobre el mismo origen. Se configuraron los "
    "servicios de autenticación (JWT), el almacenamiento de archivos para evidencias, el envío de correos "
    "y CORS restringido a los orígenes permitidos. La comunicación entre capas se verificó mediante "
    "pruebas de integración.",
)

replace_paragraph(
    "Se realizaron pruebas orientadas a verificar el funcionamiento",
    "La calidad del sistema se verificó mediante pruebas de API (Thunder Client), pruebas unitarias, de "
    "integración y de aceptación, además de tres rondas de auditoría (código, experiencia de uso y "
    "accesibilidad). En cada ronda se registraron los hallazgos, se corrigieron y se dejó constancia de "
    "la corrección en el historial de versiones. Asimismo, se verificó la compilación del backend y del "
    "frontend, la validez del esquema de la base de datos y la cobertura de las migraciones. El Cuadro "
    "2.2 resume los resultados de las auditorías realizadas.",
)

# ---------------------------------------------------------------------------
# 3) Tablas: Cuadro 2.1 (requerimientos) y Cuadro 2.2 (auditorías)
# ---------------------------------------------------------------------------
def find_table_by_header(first_cell, second_cell=None):
    for t in doc.tables:
        try:
            h0 = t.rows[0].cells[0].text.strip()
            if h0 == first_cell:
                if second_cell is None or t.rows[0].cells[1].text.strip() == second_cell:
                    return t
        except Exception:
            continue
    return None


def set_cell(cell, text, bold=False):
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    run = p.add_run(text)
    run.font.name = TNR
    run.font.size = Pt(10)
    run.font.bold = bold


def rewrite_table(table, headers, rows):
    # eliminar filas de datos (conservar la cabecera)
    trs = table._tbl.findall(qn('w:tr'))
    for tr in trs[1:]:
        table._tbl.remove(tr)
    # reescribir cabecera
    for j, h in enumerate(headers):
        set_cell(table.rows[0].cells[j], h, bold=True)
    # añadir filas nuevas
    for r in rows:
        cells = table.add_row().cells
        for j, v in enumerate(r):
            set_cell(cells[j], v, bold=False)


cuadro21 = find_table_by_header("ID")
if cuadro21 is not None:
    rewrite_table(cuadro21,
        ["ID", "Módulo", "Requerimiento", "Prioridad"],
        [
            ["RF-01", "Autenticación", "Registro del agente y su institución, inicio de sesión y activación de cuenta por invitación (JWT)", "Alta"],
            ["RF-02", "Usuarios", "Gestión del perfil y de la foto de perfil", "Media"],
            ["RF-03", "Instituciones", "Ver y editar la institución propia; catálogo de instituciones con registro completo", "Alta"],
            ["RF-04", "Docentes", "Invitar, activar, desactivar y eliminar docentes", "Alta"],
            ["RF-05", "Materias", "Registro y gestión de materias de la institución", "Alta"],
            ["RF-06", "Asignaciones", "Asignar docente a materia por periodo escolar", "Alta"],
            ["RF-07", "Solicitudes", "Crear, editar, cancelar y corregir solicitudes de clase espejo", "Alta"],
            ["RF-08", "Revisión", "Doble revisión (origen → destino) con motivo; creación automática del proyecto", "Alta"],
            ["RF-09", "Proyectos", "Gestión de proyectos de clase espejo y de sus estados", "Alta"],
            ["RF-10", "Planificación", "Planificación conjunta, confirmación de participación y reporte de planificación", "Alta"],
            ["RF-11", "Chat", "Mensajería entre los docentes del proyecto", "Media"],
            ["RF-12", "Sesiones y actividades", "Programación de sesiones y registro de actividades", "Alta"],
            ["RF-13", "Evidencias", "Carga de archivos con lista blanca de formatos (máx. 10 MB)", "Media"],
            ["RF-14", "Reportes y cierre", "Reporte de clase con firmas, evaluación final y cierre del proyecto", "Alta"],
            ["RF-15", "Recordatorios", "Agenda, calendario y caducidad automática de solicitudes sin respuesta", "Media"],
        ],
    )

cuadro22 = find_table_by_header("Módulo", "Casos de prueba")
if cuadro22 is not None:
    rewrite_table(cuadro22,
        ["Ronda", "Alcance", "Hallazgos", "Estado"],
        [
            ["1.ª", "Auditoría de código (QA)", "12 hallazgos", "Corregidos"],
            ["2.ª", "Técnica y de experiencia de uso (UX)", "3 técnicos + 6 de UX", "Corregidos"],
            ["3.ª", "Código, diseño y accesibilidad", "35 de código + 9 de diseño", "Críticos e importantes corregidos; 7 menores documentados"],
        ],
    )

# ---------------------------------------------------------------------------
# 4) Referencias: añadir documentación de los frameworks usados
# ---------------------------------------------------------------------------
# La referencia de TecNM se identifica por su texto completo (para no confundirla
# con la mención del TecNM que aparece en el marco teórico).
anchor = None
for p in doc.paragraphs:
    if p.text.strip().startswith("Tecnológico Nacional de México. (s. f.). Internacionalización"):
        anchor = p
        break
if anchor is None:
    raise ValueError("No se encontró la referencia de TecNM")

for text in [
    "Angular. (s. f.). Documentación de Angular. Recuperado el [fecha], de https://angular.dev/",
    "NestJS. (s. f.). Documentación de NestJS. Recuperado el [fecha], de https://docs.nestjs.com/",
    "Prisma. (s. f.). Documentación de Prisma ORM. Recuperado el [fecha], de https://www.prisma.io/docs",
]:
    new_el = copy.deepcopy(anchor._element)
    anchor._element.addnext(new_el)
    anchor = Paragraph(new_el, anchor._parent)
    clear_runs(anchor)
    run = anchor.add_run(text)
    run.font.name = TNR
    run.font.size = Pt(12)

doc.save(DST)

# ---------------------------------------------------------------------------
# 5) Verificación
# ---------------------------------------------------------------------------
d2 = Document(DST)
print("Guardado:", DST)
print("Párrafos:", len(d2.paragraphs), "| Tablas:", len(d2.tables))
for t in d2.tables:
    hdr = " | ".join(c.text.strip() for c in t.rows[0].cells)
    print("  tabla:", len(t.rows), "filas ->", hdr[:80])
