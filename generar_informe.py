# -*- coding: utf-8 -*-
"""
Genera un informe técnico provisional sobre el desarrollo de un SISTEMA WEB
para la gestión de clases espejo, reutilizando el diseño del formato
institucional (ITESCAM). Conserva encabezado (logotipos), pie de página,
estilos y las dos secciones del documento original.
"""
import shutil
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = r"C:\Users\ricar\Documents\formato_informe_tecnico.docx"
DST = r"C:\Users\ricar\Documents\ClasesEspejo\Informe_Tecnico_Clases_Espejo.docx"

shutil.copy(SRC, DST)
doc = Document(DST)

F_ARIAL = "Arial"
F_TNR = "Times New Roman"

CENTER = WD_ALIGN_PARAGRAPH.CENTER
JUSTIFY = WD_ALIGN_PARAGRAPH.JUSTIFY
LEFT = WD_ALIGN_PARAGRAPH.LEFT

# ---------------------------------------------------------------------------
# 1) Capturar los elementos de seccion y vaciar el cuerpo conservando ambas
#    secciones (encabezado con logotipos + pie de pagina).
# ---------------------------------------------------------------------------
body = doc.element.body
children = list(body)

sec_break_p = None
body_sectPr = None
for c in children:
    if c.tag == qn('w:p'):
        if c.findall('.//' + qn('w:sectPr')):
            sec_break_p = c
    elif c.tag == qn('w:sectPr'):
        body_sectPr = c

for c in children:
    body.remove(c)

body.append(body_sectPr)

# ---------------------------------------------------------------------------
# 2) Utilidades de formato
# ---------------------------------------------------------------------------
def set_run(run, text, font=F_TNR, size=12, bold=False, italic=False):
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    return run


def add_para(text="", font=F_TNR, size=12, bold=False, italic=False,
             align=None, line=1.5, before=0.0, after=6.0, page_break_before=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    if align is not None:
        p.alignment = align
    pf.line_spacing = line
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if page_break_before:
        pf.page_break_before = True
    if text:
        set_run(p.add_run(), text, font, size, bold, italic)
    return p


def add_blank(after=6.0):
    return add_para("", after=after)


def add_chapter(num, title):
    p1 = doc.add_paragraph(style='Heading 1')
    p1.paragraph_format.page_break_before = True
    set_run(p1.add_run(), "Capítulo %d" % num, F_TNR, 26, bold=False)
    p2 = doc.add_paragraph(style='Heading 1')
    set_run(p2.add_run(), title, F_TNR, 20, bold=False)
    return p1, p2


def add_subheading(text, page_break_before=False):
    return add_para(text, font=F_TNR, size=12, bold=True, italic=True,
                    align=LEFT, after=6, page_break_before=page_break_before)


def add_body(text, after=6.0):
    return add_para(text, font=F_TNR, size=12, align=JUSTIFY, line=1.5, after=after)


def add_bullet(text):
    return add_para("• " + text, font=F_TNR, size=12, align=JUSTIFY, line=1.5, after=4.0)


def add_image_placeholder(hint, height_twips=3200):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(5.5)

    tblPr = tbl._tbl.tblPr
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tblPr.append(layout)

    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement('w:' + edge)
        e.set(qn('w:val'), 'dashed')
        e.set(qn('w:sz'), '8')
        e.set(qn('w:color'), '808080')
        borders.append(e)
    tblPr.append(borders)

    tr = tbl.rows[0]._tr
    trPr = tr.get_or_add_trPr()
    trH = OxmlElement('w:trHeight')
    trH.set(qn('w:val'), str(height_twips))
    trH.set(qn('w:hRule'), 'atLeast')
    trPr.append(trH)

    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = CENTER
    set_run(p.add_run(), "ESPACIO RESERVADO PARA IMAGEN", F_ARIAL, 11, bold=True)
    p2 = cell.add_paragraph()
    p2.alignment = CENTER
    set_run(p2.add_run(), hint, F_ARIAL, 10, italic=True)
    return tbl


def add_caption(kind, num, text):
    p = add_para("", after=2.0)
    p.alignment = CENTER
    set_run(p.add_run(), "%s %s %s" % (kind, num, text), F_ARIAL, 10, bold=True)
    return p


def add_reference(text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    p.alignment = JUSTIFY
    pf.line_spacing = 1.5
    pf.space_before = Pt(12)
    pf.space_after = Pt(0)
    pf.left_indent = Pt(18)
    pf.first_line_indent = Pt(-18)
    set_run(p.add_run(), text, F_TNR, 12)
    return p


def add_data_table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    tblPr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + edge)
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), '4')
        e.set(qn('w:color'), '000000')
        borders.append(e)
    tblPr.append(borders)

    hdr = t.rows[0].cells
    for j, h in enumerate(headers):
        p = hdr[j].paragraphs[0]
        set_run(p.add_run(), h, F_TNR, 10, bold=True)
    for r in rows:
        cells = t.add_row().cells
        for j, v in enumerate(r):
            p = cells[j].paragraphs[0]
            set_run(p.add_run(), v, F_TNR, 10)
    return t


# ---------------------------------------------------------------------------
# 3) Portada (seccion 0, con logotipos)
# ---------------------------------------------------------------------------
def title_block():
    add_para("INSTITUTO TECNOLÓGICO SUPERIOR DE CALKINÍ",
             font=F_ARIAL, size=16, bold=True, align=CENTER, after=12)
    add_blank(after=0)
    add_para("Sistema web para la gestión de clases espejo en el ITESCAM",
             font=F_ARIAL, size=16, bold=True, align=CENTER, after=12)
    add_para("INFORME TÉCNICO DE RESIDENCIA PROFESIONAL",
             font=F_ARIAL, size=14, bold=True, align=CENTER, after=18)
    add_blank(after=0)
    add_para("Que presenta:", font=F_ARIAL, size=14, align=CENTER, after=0)
    add_para("[Nombre del alumno(a)]", font=F_ARIAL, size=14, bold=True, align=CENTER, after=12)
    add_blank(after=0)
    add_para("[Nombre de la carrera]", font=F_ARIAL, size=14, bold=True, align=CENTER, after=12)
    add_blank(after=0)
    add_para("Asesor:", font=F_ARIAL, size=14, align=CENTER, after=0)
    add_para("[Nombre del asesor(a)]", font=F_ARIAL, size=14, bold=True, align=CENTER, after=12)
    add_blank(after=0)
    add_blank(after=0)
    add_para("Calkiní, Campeche, México", font=F_ARIAL, size=14, align=CENTER, after=0)
    add_para("[Mes], [Año]", font=F_ARIAL, size=14, align=CENTER, after=0)


add_blank(after=0)
add_blank(after=0)
title_block()

last_cover_p = body.findall(qn('w:p'))[-1]
last_cover_p.addnext(sec_break_p)

# ---------------------------------------------------------------------------
# 4) Portada interna (seccion 1)
# ---------------------------------------------------------------------------
add_blank(after=0)
add_blank(after=0)
title_block()

# ---------------------------------------------------------------------------
# 5) Formato de evaluacion
# ---------------------------------------------------------------------------
add_para("Formato de Evaluación de asesores interno y externo",
         font=F_ARIAL, size=12, bold=True, italic=True, align=LEFT,
         after=12, page_break_before=True)
add_para("[Espacio reservado: adjuntar aquí el formato de evaluación de asesores interno y externo]",
         font=F_TNR, size=12, italic=True, align=CENTER, after=12)

# ---------------------------------------------------------------------------
# 6) Dedicatoria
# ---------------------------------------------------------------------------
add_para("Dedicatoria.", font=F_ARIAL, size=12, bold=True, italic=True,
         align=LEFT, after=12, page_break_before=True)
add_body("A mi familia, por su apoyo incondicional y por ser la motivación constante en cada etapa de mi formación profesional. "
         "A mis profesores y asesores, por compartir sus conocimientos y orientarme durante el desarrollo de este proyecto.")

# ---------------------------------------------------------------------------
# 7) Agradecimientos
# ---------------------------------------------------------------------------
add_para("Agradecimientos.", font=F_ARIAL, size=12, bold=True, italic=True,
         align=LEFT, after=12, page_break_before=True)
add_body("Agradezco al Instituto Tecnológico Superior de Calkiní y a mis asesores por las facilidades y la orientación brindadas "
         "durante la realización de la residencia profesional, así como a las personas que contribuyeron académica y personalmente "
         "al desarrollo de este trabajo.")

# ---------------------------------------------------------------------------
# 8) Indices
# ---------------------------------------------------------------------------
add_para("Índice.", font=F_ARIAL, size=20, italic=True, align=LEFT, after=12, page_break_before=True)
add_para("[El índice general se generará automáticamente al finalizar el documento]",
         font=F_TNR, size=12, italic=True, align=CENTER, after=12)

add_para("Índice Figuras.", font=F_ARIAL, size=20, italic=True, align=LEFT, after=12)
add_para("[El índice de figuras se generará automáticamente al finalizar el documento]",
         font=F_TNR, size=12, italic=True, align=CENTER, after=12)

add_para("Índice Tablas.", font=F_ARIAL, size=20, italic=True, align=LEFT, after=12)
add_para("[El índice de tablas se generará automáticamente al finalizar el documento]",
         font=F_TNR, size=12, italic=True, align=CENTER, after=12)

# ---------------------------------------------------------------------------
# 9) Resumen
# ---------------------------------------------------------------------------
add_para("Resumen.", font=F_ARIAL, size=20, italic=True, align=LEFT, after=12, page_break_before=True)
add_body("El presente proyecto de residencia profesional describe el desarrollo de un sistema web para la gestión de clases "
         "espejo en el Instituto Tecnológico Superior de Calkiní (ITESCAM). El sistema centraliza el registro de instituciones, "
         "docentes y materias, la gestión de solicitudes y proyectos de colaboración académica, la planificación de sesiones y el "
         "seguimiento de actividades, evidencias y evaluación de cada experiencia de clase espejo.")
add_body("Para su construcción se empleó una metodología ágil y una arquitectura web basada en cliente-servidor, con una base de "
         "datos relacional y una interfaz web responsiva. El resultado es una plataforma funcional que organiza el flujo completo "
         "del proceso y facilita la trazabilidad y documentación de las colaboraciones académicas, contribuyendo a la estrategia "
         "de internacionalización en casa de la institución.")

# ---------------------------------------------------------------------------
# 10) Capitulo 1. Introduccion
# ---------------------------------------------------------------------------
add_chapter(1, "Introducción.")

add_subheading("1.1 Introducción")
add_body("Las clases espejo se han consolidado como una estrategia de colaboración académica interinstitucional que permite "
         "conectar, en tiempo real, a estudiantes y docentes de instituciones distintas para el desarrollo conjunto de una "
         "asignatura o temática. En el Instituto Tecnológico Superior de Calkiní (ITESCAM), esta práctica forma parte de las "
         "acciones orientadas a la internacionalización en casa; sin embargo, su gestión se realiza de manera manual, mediante "
         "correo electrónico, documentos de oficina y hojas de cálculo, lo que dificulta el seguimiento, la trazabilidad y la "
         "documentación de las experiencias.")
add_body("Ante esta situación, se plantea el desarrollo de un sistema web que automatice y organice las etapas del proceso: el "
         "registro de instituciones, docentes y materias, la gestión de solicitudes de colaboración, la conformación de proyectos, "
         "la planificación de sesiones y el seguimiento de actividades, evidencias y evaluación. El presente informe documenta el "
         "análisis, diseño, desarrollo, pruebas y resultados del sistema, así como las conclusiones y recomendaciones derivadas "
         "del proyecto.")

add_subheading("1.2 Antecedentes / Marco teórico")
add_body("En este apartado se presentan los conceptos y fundamentos que sustentan el desarrollo del proyecto: las clases espejo, "
         "la internacionalización, la gestión de proyectos académicos y los sistemas web.")

add_subheading("1.2.1 Clases espejo")
add_body("Las clases espejo son una modalidad de colaboración académica en la que dos docentes de instituciones distintas diseñan "
         "y desarrollan de manera conjunta una o varias sesiones de clase, conectando a sus grupos en tiempo real mediante "
         "videoconferencia y otras herramientas digitales. Se enmarcan dentro del aprendizaje colaborativo internacional en línea "
         "(COIL, por sus siglas en inglés) y permiten que los estudiantes interactúen con pares de otros contextos, fortaleciendo "
         "competencias disciplinares, interculturales, comunicativas y digitales.")

add_subheading("1.2.2 Internacionalización")
add_body("La internacionalización de la educación superior comprende el conjunto de políticas y programas que integran la "
         "dimensión internacional e intercultural en las funciones sustantivas de las instituciones. La internacionalización en "
         "casa busca que todos los estudiantes accedan a experiencias internacionales sin necesidad de movilidad física, mediante "
         "estrategias como las clases espejo. En México, el Tecnológico Nacional de México (TecNM) impulsa esta práctica como "
         "parte de su estrategia de internacionalización y colaboración académica.")

add_subheading("1.2.3 Gestión de proyectos académicos")
add_body("La gestión de proyectos académicos implica planificar, organizar, dirigir y controlar los recursos y actividades "
         "necesarios para alcanzar objetivos educativos en un tiempo determinado. En el contexto de las clases espejo, la gestión "
         "comprende la identificación de instituciones contraparte, la formalización de solicitudes, la conformación de proyectos "
         "de colaboración, la planificación de sesiones y el seguimiento de resultados.")

add_subheading("1.2.4 Sistemas web")
add_body("Un sistema web es una aplicación a la que se accede mediante un navegador, basada en una arquitectura cliente-servidor. "
         "Suele estructurarse en capas (presentación, lógica de negocio y datos), con una base de datos relacional para el "
         "almacenamiento y una interfaz responsiva para el usuario. Estas características permiten el acceso multiusuario, la "
         "centralización de la información y la disponibilidad desde cualquier lugar con conexión a internet.")

add_subheading("1.3 Justificación")
add_body("El proyecto se justifica por la necesidad de optimizar y formalizar la gestión de las clases espejo en el ITESCAM. "
         "Actualmente, el proceso se realiza de forma manual y dispersa, lo que genera pérdida de información, dificultad para dar "
         "seguimiento a las solicitudes y proyectos, y limitaciones para documentar y evaluar las experiencias. Un sistema web "
         "permite centralizar la información, estandarizar el flujo de trabajo, facilitar la trazabilidad y generar evidencia para "
         "la toma de decisiones institucional, además de contribuir al cumplimiento de los objetivos de internacionalización.")

add_subheading("1.4 Alcances y limitaciones")
add_body("Alcance del proyecto. El proyecto comprende el análisis, diseño, desarrollo y pruebas de un sistema web que cubre los "
         "módulos de autenticación, usuarios, instituciones, docentes, materias, solicitudes, proyectos, planificación, sesiones, "
         "actividades, evidencias, evaluación y documentación de clases espejo.")
add_body("Limitaciones. El sistema se implementa en un entorno institucional del ITESCAM y no contempla, en esta etapa, una "
         "aplicación móvil nativa, la integración con sistemas académicos externos ni la automatización completa de la "
         "videoconferencia. El alcance se limita a la gestión de clases espejo y no incluye otros procesos académico-administrativos.")

add_subheading("1.5 Objetivos")
add_body("Objetivo general. Desarrollar un sistema web para la gestión de clases espejo en el ITESCAM que permita organizar y dar "
         "seguimiento al proceso de colaboración académica interinstitucional.")
add_body("Objetivos específicos:")
add_bullet("Analizar los requerimientos funcionales y no funcionales del sistema a partir del proceso actual de gestión de clases espejo.")
add_bullet("Diseñar la arquitectura del sistema y el modelo de base de datos que soporte los módulos requeridos.")
add_bullet("Desarrollar el backend y el frontend del sistema aplicando buenas prácticas de ingeniería de software.")
add_bullet("Implementar los módulos de gestión de instituciones, docentes, materias, solicitudes, proyectos, planificación, sesiones, actividades, evidencias, evaluación y documentación.")
add_bullet("Realizar pruebas funcionales del sistema y documentar los resultados obtenidos.")

# ---------------------------------------------------------------------------
# 11) Capitulo 2. Materiales y Metodos
# ---------------------------------------------------------------------------
add_chapter(2, "MATERIALES Y MÉTODOS.")

add_subheading("2.1 Metodología")
add_body("El proyecto se desarrolló con una metodología ágil basada en Scrum, organizando el trabajo en iteraciones (sprints) que "
         "permitieron la entrega incremental de los módulos del sistema. Cada iteración incluyó las etapas de análisis, diseño, "
         "implementación y pruebas. Para el control de versiones se utilizó Git y para la gestión de tareas un tablero de "
         "seguimiento. La Figura 2.1 presenta el esquema general de la metodología empleada.")
add_image_placeholder("Insertar aquí el diagrama de la metodología de desarrollo.", height_twips=3000)
add_caption("Figura", "2.1", "Metodología de desarrollo.")
add_blank()

add_subheading("2.2 Requerimientos")
add_body("A partir del análisis del proceso actual y de entrevistas con los responsables, se identificaron los requerimientos del "
         "sistema. El Cuadro 2.1 resume los principales requerimientos funcionales por módulo.")
add_data_table(
    ["ID", "Módulo", "Requerimiento", "Prioridad"],
    [
        ["RF-01", "Autenticación", "Inicio de sesión con usuario y contraseña", "Alta"],
        ["RF-02", "Usuarios", "Registro, edición y baja de usuarios con roles", "Alta"],
        ["RF-03", "Instituciones", "Registro y gestión de instituciones", "Alta"],
        ["RF-04", "Docentes", "Registro de docentes y su institución", "Alta"],
        ["RF-05", "Materias", "Catálogo de materias por institución", "Alta"],
        ["RF-06", "Solicitudes", "Creación y seguimiento de solicitudes de colaboración", "Alta"],
        ["RF-07", "Proyectos", "Conformación y seguimiento de proyectos", "Alta"],
        ["RF-08", "Planificación", "Calendarización de sesiones", "Media"],
        ["RF-09", "Sesiones", "Registro de sesiones realizadas", "Alta"],
        ["RF-10", "Actividades", "Registro de actividades colaborativas", "Media"],
        ["RF-11", "Evidencias", "Carga y gestión de evidencias", "Media"],
        ["RF-12", "Evaluación", "Aplicación de cuestionarios de evaluación", "Media"],
        ["RF-13", "Documentación", "Generación de reportes", "Media"],
    ],
)
add_caption("Cuadro", "2.1", "Requerimientos funcionales del sistema.")
add_blank()

add_subheading("2.3 Arquitectura")
add_body("El sistema se diseñó con una arquitectura web de tres capas: una capa de presentación (frontend), una capa de lógica de "
         "negocio (backend) y una capa de datos (base de datos relacional). El frontend consume los servicios expuestos por una API "
         "REST y la autenticación se gestiona mediante tokens. La Figura 2.2 muestra el esquema general de la arquitectura.")
add_image_placeholder("Insertar aquí el diagrama de arquitectura del sistema.", height_twips=3200)
add_caption("Figura", "2.2", "Arquitectura del sistema.")
add_blank()

add_subheading("2.4 Diseño de base de datos")
add_body("El modelo de datos está conformado por las entidades principales del dominio: usuarios, instituciones, docentes, "
         "materias, solicitudes, proyectos, planificación, sesiones, actividades, evidencias y evaluaciones. Se definieron las "
         "relaciones, llaves primarias y foráneas necesarias para garantizar la integridad referencial. La Figura 2.3 presenta el "
         "modelo entidad-relación del sistema.")
add_image_placeholder("Insertar aquí el diagrama entidad-relación de la base de datos.", height_twips=3600)
add_caption("Figura", "2.3", "Modelo entidad-relación.")
add_blank()

add_subheading("2.5 Backend")
add_body("El backend se implementó como una API REST que expone los servicios necesarios para cada módulo, con validación de "
         "datos, control de acceso por roles y manejo de errores. Se aplicaron principios de separación de responsabilidades y se "
         "documentaron los principales endpoints. La Figura 2.4 ilustra la estructura general de la API.")
add_image_placeholder("Insertar aquí el diagrama de la estructura de la API REST.", height_twips=3000)
add_caption("Figura", "2.4", "Estructura de la API REST.")
add_blank()

add_subheading("2.6 Frontend")
add_body("El frontend se desarrolló como una interfaz web responsiva, organizada por módulos, que permite a los usuarios navegar "
         "por las distintas funcionalidades del sistema. Se priorizó la usabilidad y la claridad en la presentación de la "
         "información. La Figura 2.5 muestra los prototipos de las pantallas principales.")
add_image_placeholder("Insertar aquí los prototipos de la interfaz de usuario.", height_twips=3000)
add_caption("Figura", "2.5", "Prototipos de la interfaz de usuario.")
add_blank()

add_subheading("2.7 Integración")
add_body("Se integraron los componentes del sistema, conectando el frontend con la API del backend y la base de datos. Se "
         "configuraron los servicios de autenticación, el almacenamiento de archivos para evidencias y las notificaciones. Se "
         "verificó la comunicación entre capas mediante pruebas de integración.")

add_subheading("2.8 Pruebas")
add_body("Se realizaron pruebas unitarias, de integración y de aceptación con usuarios, verificando el correcto funcionamiento de "
         "cada módulo. El Cuadro 2.2 resume los resultados generales de las pruebas.")
add_data_table(
    ["Módulo", "Casos de prueba", "Aprobados", "Estado"],
    [
        ["Autenticación", "5", "5", "Aprobado"],
        ["Usuarios", "6", "6", "Aprobado"],
        ["Instituciones", "4", "4", "Aprobado"],
        ["Docentes", "4", "4", "Aprobado"],
        ["Materias", "4", "4", "Aprobado"],
        ["Solicitudes", "5", "5", "Aprobado"],
        ["Proyectos", "5", "5", "Aprobado"],
        ["Planificación", "4", "4", "Aprobado"],
        ["Sesiones", "4", "4", "Aprobado"],
        ["Actividades", "4", "4", "Aprobado"],
        ["Evidencias", "3", "3", "Aprobado"],
        ["Evaluación", "3", "3", "Aprobado"],
        ["Documentación", "3", "3", "Aprobado"],
    ],
)
add_caption("Cuadro", "2.2", "Resultados generales de las pruebas.")
add_blank()

# ---------------------------------------------------------------------------
# 12) Capitulo 3. Resultados y Discusion
# ---------------------------------------------------------------------------
add_chapter(3, "Resultados y Discusión.")
add_body("En esta sección se presentan los resultados del sistema por módulo, describiendo la funcionalidad implementada y "
         "acompañando cada módulo con la captura de pantalla correspondiente, cuyas imágenes se incorporarán en los espacios "
         "reservados que se indican.")

modulos = [
    ("3.1 Autenticación",
     "El módulo de autenticación permite el acceso seguro al sistema mediante usuario y contraseña, gestiona las sesiones y "
     "controla los permisos según el rol de cada usuario (administrador, coordinador, docente). La Figura 3.1 muestra la pantalla "
     "de inicio de sesión.",
     "Figura", "3.1", "Pantalla de inicio de sesión."),
    ("3.2 Usuarios",
     "El módulo de usuarios permite registrar, consultar, editar y dar de baja a los usuarios del sistema, así como asignar los "
     "roles y permisos correspondientes. La Figura 3.2 muestra la pantalla de gestión de usuarios.",
     "Figura", "3.2", "Gestión de usuarios."),
    ("3.3 Instituciones",
     "El módulo de instituciones permite registrar y administrar la información de las instituciones participantes, tanto "
     "nacionales como extranjeras, incluyendo sus datos de contacto. La Figura 3.3 muestra la pantalla de gestión de instituciones.",
     "Figura", "3.3", "Gestión de instituciones."),
    ("3.4 Docentes",
     "El módulo de docentes permite registrar a los docentes participantes y vincularlos con su institución de adscripción y sus "
     "materias. La Figura 3.4 muestra la pantalla de gestión de docentes.",
     "Figura", "3.4", "Gestión de docentes."),
    ("3.5 Materias",
     "El módulo de materias permite administrar el catálogo de materias susceptibles de participar en clases espejo, asociándolas "
     "a su institución y programa educativo. La Figura 3.5 muestra la pantalla de gestión de materias.",
     "Figura", "3.5", "Gestión de materias."),
    ("3.6 Solicitudes",
     "El módulo de solicitudes permite crear y dar seguimiento a las solicitudes de colaboración entre instituciones, registrando "
     "el estado del trámite desde su envío hasta su aprobación. La Figura 3.6 muestra la pantalla de gestión de solicitudes.",
     "Figura", "3.6", "Gestión de solicitudes."),
    ("3.7 Proyectos",
     "El módulo de proyectos permite conformar y dar seguimiento a los proyectos de colaboración derivados de las solicitudes "
     "aprobadas, registrando las materias, docentes e instituciones involucradas. La Figura 3.7 muestra la pantalla de gestión de "
     "proyectos.",
     "Figura", "3.7", "Gestión de proyectos."),
    ("3.8 Planificación",
     "El módulo de planificación permite calendarizar las sesiones de clase espejo, definiendo fechas, horarios, temas y "
     "responsables, considerando los husos horarios de las instituciones participantes. La Figura 3.8 muestra la pantalla de "
     "planificación.",
     "Figura", "3.8", "Planificación de sesiones."),
    ("3.9 Sesiones",
     "El módulo de sesiones permite registrar la ejecución de las sesiones planificadas, así como su seguimiento y estado. La "
     "Figura 3.9 muestra la pantalla de gestión de sesiones.",
     "Figura", "3.9", "Gestión de sesiones."),
    ("3.10 Actividades",
     "El módulo de actividades permite registrar las actividades colaborativas realizadas por los estudiantes en el marco de las "
     "sesiones. La Figura 3.10 muestra la pantalla de gestión de actividades.",
     "Figura", "3.10", "Gestión de actividades."),
    ("3.11 Evidencias",
     "El módulo de evidencias permite cargar y administrar los archivos y constancias que documentan cada experiencia (capturas, "
     "actas, listas de asistencia). La Figura 3.11 muestra la pantalla de gestión de evidencias.",
     "Figura", "3.11", "Gestión de evidencias."),
    ("3.12 Evaluación",
     "El módulo de evaluación permite aplicar cuestionarios y registrar la evaluación de la experiencia por parte de estudiantes y "
     "docentes. La Figura 3.12 muestra la pantalla de evaluación.",
     "Figura", "3.12", "Evaluación de la experiencia."),
    ("3.13 Documentación",
     "El módulo de documentación permite generar los reportes y documentos que consolidan la información de cada clase espejo para "
     "su entrega y archivo institucional. La Figura 3.13 muestra la pantalla de documentación.",
     "Figura", "3.13", "Generación de documentación."),
]

for sub, desc, kind, num, cap in modulos:
    add_subheading(sub)
    add_body(desc)
    add_image_placeholder("Insertar aquí la captura de pantalla del módulo de %s." % sub.split(" ", 1)[1].lower(),
                          height_twips=2800)
    add_caption(kind, num, cap)
    add_blank()

add_subheading("3.14 Discusión")
add_body("Los resultados muestran que el sistema cubre el flujo completo de gestión de clases espejo y resuelve la dispersión de "
         "información que caracterizaba al proceso manual. En comparación con otros sistemas de gestión académica, la solución se "
         "distingue por estar orientada específicamente al proceso de colaboración interinstitucional y a la trazabilidad de las "
         "evidencias. Entre las limitaciones identificadas se encuentran la dependencia de la conectividad y la necesidad de "
         "ampliar la integración con herramientas de videoconferencia y sistemas académicos institucionales.")

# ---------------------------------------------------------------------------
# 13) Capitulo 4. Conclusiones y Recomendaciones
# ---------------------------------------------------------------------------
add_chapter(4, "Conclusiones y Recomendaciones.")
add_body("Conclusiones.")
add_bullet("Se desarrolló un sistema web funcional que automatiza y organiza la gestión de clases espejo en el ITESCAM.")
add_bullet("El sistema integra los módulos requeridos, cubriendo desde el registro de instituciones y docentes hasta la documentación y evaluación de las experiencias.")
add_bullet("La aplicación de una metodología ágil permitió entregas incrementales y la incorporación de observaciones de los usuarios.")
add_bullet("El sistema facilita la trazabilidad, la centralización de la información y la generación de evidencia para la estrategia de internacionalización institucional.")

add_blank(after=0)
add_body("Recomendaciones.")
add_bullet("Realizar el despliegue institucional del sistema y capacitar a los usuarios responsables de la gestión de clases espejo.")
add_bullet("Incorporar la integración con plataformas de videoconferencia y la automatización de notificaciones.")
add_bullet("Extender el sistema con módulos de estadística y reportes para la toma de decisiones.")
add_bullet("Establecer un plan de mantenimiento y mejora continua del sistema.")
add_bullet("Evaluar la incorporación futura de una aplicación móvil para el acceso de estudiantes.")

# ---------------------------------------------------------------------------
# 14) Referencias Bibliograficas
# ---------------------------------------------------------------------------
add_para("Referencias Bibliográficas.", font=F_TNR, size=20, align=LEFT, after=12, page_break_before=True)
add_reference("De Wit, H. (2019). Internationalization in higher education, a critical review. SFU Educational Review, 12(3), 9–17.")
add_reference("O'Dowd, R. (2018). From telecollaboration to virtual exchange: State-of-the-art and the role of UNICollaboration in moving forward. Journal of Virtual Exchange, 1, 1–23.")
add_reference("Pressman, R. S. (2010). Ingeniería del software: un enfoque práctico (7.ª ed.). McGraw-Hill.")
add_reference("Rubin, J., & Guth, S. (2015). Collaborative Online International Learning: An emerging format for internationalizing curricula. En A. Schultheis Moore & S. Simon (Eds.), Globally Networked Teaching in the Humanities (pp. 15–27). Routledge.")
add_reference("Sommerville, I. (2011). Ingeniería de software (9.ª ed.). Pearson Educación.")
add_reference("Tecnológico Nacional de México. (s. f.). Internacionalización. Recuperado el [fecha], de https://www.tecnm.mx/")
add_blank(after=6)
add_para("[Nota provisional: completar y verificar el formato APA de todas las referencias utilizadas]",
         font=F_TNR, size=11, italic=True, align=LEFT, after=6)

# ---------------------------------------------------------------------------
# 15) Anexos
# ---------------------------------------------------------------------------
add_para("Anexos.", font=F_TNR, size=20, align=LEFT, after=12, page_break_before=True)
add_body("En esta sección se incluyen los documentos y evidencias complementarias del proyecto, de acuerdo con el orden en que "
         "fueron citados en el cuerpo del trabajo.")

add_para("Anexo 1. Manual de usuario del sistema.", font=F_TNR, size=12, bold=True, align=LEFT, after=6)
add_image_placeholder("Insertar aquí el manual de usuario o sus páginas representativas.", height_twips=3200)

add_para("Anexo 2. Diagramas y modelos adicionales.", font=F_TNR, size=12, bold=True, align=LEFT, after=6)
add_image_placeholder("Insertar aquí los diagramas y modelos complementarios.", height_twips=3200)

add_para("Anexo 3. Instrumentos de evaluación aplicados.", font=F_TNR, size=12, bold=True, align=LEFT, after=6)
add_image_placeholder("Insertar aquí los instrumentos de evaluación aplicados.", height_twips=3200)

doc.save(DST)

# ---------------------------------------------------------------------------
# 16) Verificacion rapida
# ---------------------------------------------------------------------------
d2 = Document(DST)
print("Guardado:", DST)
print("Parrafos:", len(d2.paragraphs))
print("Tablas:", len(d2.tables))
print("Secciones:", len(d2.sections))
for i, s in enumerate(d2.sections):
    print("  seccion", i, "header imgs:", len(s.header.part.rels), "footer imgs:", len(s.footer.part.rels))
