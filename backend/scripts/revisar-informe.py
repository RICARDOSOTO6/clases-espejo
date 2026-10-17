"""Revisa el informe tecnico mas reciente: cuantos parrafos, tablas y si las cifras
y los temas nuevos aparecen. Uso: python revisar-informe.py <archivo.docx>"""
import sys, re
from docx import Document

ruta = sys.argv[1]
doc = Document(ruta)

parrafos = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
tablas = doc.tables
print(f"archivo         : {ruta}")
print(f"parrafos        : {len(parrafos)}")
print(f"tablas          : {len(tablas)}")
print(f"secciones       : {len(doc.sections)}")
print()

# Temas que deberian aparecer si el informe esta al dia
temas = {
    "CACE / presidente": r"CACE|presidencia|presidente",
    "Secretaria": r"secretar",
    "Capacidades": r"capacidad(es)?\b|CapacidadesGuard",
    "Certificados": r"certificad",
    "Verificacion publica": r"verificaci[oó]n p[uú]blica|/verificar|folio",
    "QR": r"\bQR\b",
    "Clases espejo (listas)": r"lista de alumnos|alumnos",
    "Instituciones aprobadas": r"APROBADA|PENDIENTE|vetad",
    "Recordatorios": r"recordatorio",
    "Calendario .ics": r"\.ics|calendario",
    "Auditorias": r"auditor[ií]a",
    "Roles": r"\brol\b|roles",
}
texto = "\n".join(parrafos)
for nombre, patron in temas.items():
    n = len(re.findall(patron, texto, re.IGNORECASE))
    print(f"{'SI' if n else 'NO':>3}  {nombre:<26} {n:>4} menciones")

print()
print("=== frases con cifras (para comprobar si estan al dia) ===")
patron_cifra = re.compile(
    r"[^.]*\b(\d{1,3})\s*(tablas|modelos|migraciones|endpoints|rutas|controladores|"
    r"servicios|m[oó]dulos|roles|capacidades|casos de uso|pantallas)\b[^.]*\.",
    re.IGNORECASE,
)
vistos = set()
for m in patron_cifra.finditer(texto):
    frase = " ".join(m.group(0).split())
    if frase[:70] in vistos:
        continue
    vistos.add(frase[:70])
    print(" -", frase[:220])

print()
print("=== encabezados de las tablas (primera fila) ===")
for i, t in enumerate(tablas[:25], 1):
    try:
        celdas = [c.text.strip()[:28] for c in t.rows[0].cells]
        print(f" {i:>2}. ({len(t.rows)}x{len(t.columns)}) {' | '.join(celdas[:6])}")
    except Exception as e:
        print(f" {i:>2}. (no se pudo leer: {e})")
